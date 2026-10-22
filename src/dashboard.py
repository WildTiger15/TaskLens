"""Build the time study dashboard: an Excel workbook with native charts, plus
clean CSV tables that load straight into Power BI Desktop.

Reads (written by earlier steps):
  outputs/<model>/run_metrics.csv   from src/segment.py
  outputs/<model>/segments.csv      from src/segment.py
  outputs/metrics.csv               from src/evaluate.py (optional)
Writes:
  outputs/tasklens_dashboard.xlsx   sheets: Runs, Steps, Variation, Evaluation, About
  outputs/powerbi/runs.csv, step_times.csv, steps.csv, evaluation.csv

Usage:
  python src/dashboard.py
  python src/dashboard.py --model claude-haiku-5-5
"""

import argparse
import json
import sys
from datetime import date
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.chart.marker import DataPoint
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import config  # noqa: E402

STEPS = json.loads(config.STEPS_FILE.read_text(encoding="utf-8"))
NOT_BOTTLENECK = {"idle", "unclear"}   # waiting is not a process step to fix
BOLD = Font(bold=True)
HEADER_FILL = PatternFill("solid", fgColor="DDE6F3")
FLAG_FILL = PatternFill("solid", fgColor="FFE08A")


# ---------- tables (these become both the Excel sheets and the Power BI CSVs) ----------

def load_inputs(model, outputs_dir):
    model_dir = Path(outputs_dir) / model
    runs_path = model_dir / "run_metrics.csv"
    segments_path = model_dir / "segments.csv"
    if not runs_path.exists() or not segments_path.exists():
        raise SystemExit(f"Missing {runs_path} or {segments_path}. Run src/segment.py first.")
    runs = pd.read_csv(runs_path)
    segments = pd.read_csv(segments_path)
    eval_path = Path(outputs_dir) / "metrics.csv"
    evaluation = pd.read_csv(eval_path) if eval_path.exists() else None
    return runs, segments, evaluation


def step_times_table(segments, runs):
    """Long table: one row per run and step with the total seconds in that step (0 if absent)."""
    totals = segments.groupby(["run", "step"])["duration_s"].sum()
    rows = [{"run": run, "step": step, "seconds": round(float(totals.get((run, step), 0.0)), 3)}
            for run in runs["run"] for step in STEPS]
    return pd.DataFrame(rows)


def steps_table(step_times):
    """Mean, standard deviation, min and max seconds per run for each step, and the bottleneck flag."""
    table = step_times.groupby("step", sort=False)["seconds"].agg(
        n_runs="count",
        runs_with_step=lambda s: int((s > 0).sum()),
        mean_s="mean",
        sd_s=lambda s: s.std(ddof=1) if len(s) > 1 else 0.0,
        min_s="min",
        max_s="max",
    ).reset_index()
    table[["mean_s", "sd_s", "min_s", "max_s"]] = table[["mean_s", "sd_s", "min_s", "max_s"]].round(2)
    candidates = table[~table["step"].isin(NOT_BOTTLENECK)]
    bottleneck = candidates.loc[candidates["mean_s"].idxmax(), "step"] if len(candidates) else ""
    table["is_bottleneck"] = table["step"] == bottleneck
    return table


def runs_table(runs):
    columns = ["run", "cycle_s", "touch_s", "wait_s", "unclear_s", "rework_s",
               "fix_error_count", "fix_error_s", "repeat_step_count", "repeat_step_s"]
    return runs[columns].copy()


def variation_stats(runs):
    cycle = runs["cycle_s"]
    mean = cycle.mean()
    sd = cycle.std(ddof=1) if len(cycle) > 1 else 0.0
    return {"n_runs": len(cycle), "mean_s": mean, "sd_s": sd, "min_s": cycle.min(),
            "max_s": cycle.max(), "range_s": cycle.max() - cycle.min(),
            "cv_pct": 100 * sd / mean if mean else 0.0}


def evaluation_summary(evaluation):
    """One row per model: n and the mean of each evaluation metric."""
    columns = ["frame_accuracy_raw", "frame_accuracy", "cycle_abs_error_s", "cycle_pct_error",
               "step_duration_mae_s", "order_distance", "cost_usd"]
    data = evaluation.copy()
    for column in columns:
        data[column] = pd.to_numeric(data[column], errors="coerce")
    summary = data.groupby("model")[columns].mean().round(4).reset_index()
    summary.insert(1, "n_runs", data.groupby("model").size().values)
    return summary


# ---------- Excel helpers ----------

def write_table(ws, df, start_row=1, start_col=1):
    """Write a DataFrame with a styled header. Returns the last row written."""
    for j, name in enumerate(df.columns):
        cell = ws.cell(row=start_row, column=start_col + j, value=name)
        cell.font, cell.fill = BOLD, HEADER_FILL
    for i, row in enumerate(df.itertuples(index=False), start=1):
        for j, value in enumerate(row):
            if hasattr(value, "item"):  # numpy number -> plain Python number
                value = value.item()
            cell = ws.cell(row=start_row + i, column=start_col + j, value=value)
            if isinstance(value, float):
                cell.number_format = "0.0##"
    for j, name in enumerate(df.columns):
        ws.column_dimensions[get_column_letter(start_col + j)].width = max(11, len(str(name)) + 3)
    return start_row + len(df)


def add_note(ws, cell, text):
    ws[cell] = text
    ws[cell].font = Font(italic=True, color="555555")


def bar_chart(title, y_title, data_ref, cats_ref, stacked=False):
    chart = BarChart()
    chart.type = "col"
    chart.title, chart.y_axis.title = title, y_title
    chart.add_data(data_ref, titles_from_data=True)
    chart.set_categories(cats_ref)
    if stacked:
        chart.grouping, chart.overlap = "stacked", 100
    chart.height, chart.width = 8, 16
    return chart


# ---------- sheets ----------

def build_runs_sheet(wb, runs, model):
    ws = wb.active
    ws.title = "Runs"
    last = write_table(ws, runs)
    n = len(runs)
    # Touch + wait + unclear stack up to the cycle time of each run
    chart = bar_chart(f"Cycle time split per run (n = {n})", "seconds",
                      Reference(ws, min_col=3, max_col=5, min_row=1, max_row=last),
                      Reference(ws, min_col=1, min_row=2, max_row=last), stacked=True)
    ws.add_chart(chart, "L2")
    add_note(ws, f"A{last + 2}", f"Source: outputs/{model}/run_metrics.csv (also outputs/powerbi/runs.csv). "
                                 "Touch + wait + unclear = cycle time.")


def build_steps_sheet(wb, steps, model):
    ws = wb.create_sheet("Steps")
    last = write_table(ws, steps)
    flag_col = list(steps.columns).index("is_bottleneck") + 1
    for row in range(2, last + 1):
        if ws.cell(row=row, column=flag_col).value is True:
            for col in range(1, len(steps.columns) + 1):
                ws.cell(row=row, column=col).fill = FLAG_FILL

    chart = bar_chart(f"Mean seconds per run in each step (n = {int(steps['n_runs'].max())} runs)", "seconds",
                      Reference(ws, min_col=4, min_row=1, max_row=last),
                      Reference(ws, min_col=1, min_row=2, max_row=last))
    chart.legend = None
    # Color the bottleneck bar so it stands out in the chart too
    for index, is_bottleneck in enumerate(steps["is_bottleneck"]):
        if is_bottleneck:
            point = DataPoint(idx=index)
            point.graphicalProperties.solidFill = "E0A100"
            chart.series[0].dPt.append(point)
    ws.add_chart(chart, "K2")

    flagged = steps[steps["is_bottleneck"]]
    if len(flagged):
        b = flagged.iloc[0]
        add_note(ws, f"A{last + 2}", f"Bottleneck (highlighted): {b['step']}, mean {b['mean_s']:.1f} s per run, "
                                     f"SD {b['sd_s']:.1f} s, n = {int(b['n_runs'])} runs. idle and unclear are not counted.")
    add_note(ws, f"A{last + 3}", f"Source: outputs/{model}/segments.csv, summed per run "
                                 "(outputs/powerbi/step_times.csv and steps.csv). Steps missing from a run count as 0 s.")


def build_variation_sheet(wb, runs):
    ws = wb.create_sheet("Variation")
    last = write_table(ws, runs[["run", "cycle_s"]])
    stats = variation_stats(runs)

    chart = LineChart()
    chart.title, chart.y_axis.title = f"Cycle time per run (n = {stats['n_runs']})", "seconds"
    chart.add_data(Reference(ws, min_col=2, min_row=1, max_row=last), titles_from_data=True)
    chart.set_categories(Reference(ws, min_col=1, min_row=2, max_row=last))
    chart.legend = None
    chart.height, chart.width = 8, 16
    ws.add_chart(chart, "E2")

    stat_rows = pd.DataFrame([{"statistic": name, "value": round(float(value), 2)} for name, value in stats.items()])
    stats_last = write_table(ws, stat_rows, start_row=last + 2)
    add_note(ws, f"A{stats_last + 2}",
             f"Spread: cycle time ran from {stats['min_s']:.1f} to {stats['max_s']:.1f} s across {stats['n_runs']} runs "
             f"(mean {stats['mean_s']:.1f} s, SD {stats['sd_s']:.1f} s, coefficient of variation {stats['cv_pct']:.0f}%).")


def build_evaluation_sheet(wb, evaluation):
    ws = wb.create_sheet("Evaluation")
    if evaluation is None or evaluation.empty:
        add_note(ws, "A1", "No evaluation yet. Run src/evaluate.py, then rebuild the dashboard.")
        return
    summary = evaluation_summary(evaluation)
    last = write_table(ws, summary)
    add_note(ws, f"A{last + 1}", "Means over runs. Accuracy is a share from 0 to 1. Source: outputs/metrics.csv.")
    ws.cell(row=last + 3, column=1, value="Per run").font = BOLD
    write_table(ws, evaluation, start_row=last + 4)


def build_about_sheet(wb, model, runs):
    ws = wb.create_sheet("About")
    lines = [
        ("TaskLens time study dashboard", BOLD),
        (f"Model used for the time study: {model}", None),
        (f"Runs: n = {len(runs)}. This is a small study; read every number with its n.", None),
        (f"Built: {date.today().isoformat()} by src/dashboard.py", None),
        ("Every number comes from the CSV files in outputs/. Power BI versions are in outputs/powerbi/.", None),
        ("All data is dummy invoice data; recordings show only the author's own screen.", None),
    ]
    for i, (text, font) in enumerate(lines, start=1):
        ws.cell(row=i, column=1, value=text).font = font or Font()
    ws.column_dimensions["A"].width = 100


# ---------- main ----------

def build_dashboard(model=config.MODEL, outputs_dir=config.OUTPUTS_DIR):
    runs_raw, segments, evaluation = load_inputs(model, outputs_dir)
    runs = runs_table(runs_raw)
    step_times = step_times_table(segments, runs)
    steps = steps_table(step_times)

    powerbi_dir = Path(outputs_dir) / "powerbi"
    powerbi_dir.mkdir(parents=True, exist_ok=True)
    runs.to_csv(powerbi_dir / "runs.csv", index=False)
    step_times.to_csv(powerbi_dir / "step_times.csv", index=False)
    steps.to_csv(powerbi_dir / "steps.csv", index=False)
    if evaluation is not None:
        evaluation.to_csv(powerbi_dir / "evaluation.csv", index=False)

    wb = Workbook()
    build_runs_sheet(wb, runs, model)
    build_steps_sheet(wb, steps, model)
    build_variation_sheet(wb, runs)
    build_evaluation_sheet(wb, evaluation)
    build_about_sheet(wb, model, runs)
    workbook_path = Path(outputs_dir) / "tasklens_dashboard.xlsx"
    wb.save(workbook_path)
    return workbook_path, powerbi_dir, steps, runs


def main():
    parser = argparse.ArgumentParser(description="Build the Excel dashboard and Power BI CSVs.")
    parser.add_argument("--model", default=config.MODEL)
    args = parser.parse_args()

    workbook_path, powerbi_dir, steps, runs = build_dashboard(args.model)
    flagged = steps[steps["is_bottleneck"]]
    if len(flagged):
        b = flagged.iloc[0]
        print(f"Bottleneck: {b['step']} at {b['mean_s']:.1f} s mean per run (n = {len(runs)} runs)")
    print(f"Wrote {workbook_path}")
    print(f"Wrote Power BI tables to {powerbi_dir}")


if __name__ == "__main__":
    main()
