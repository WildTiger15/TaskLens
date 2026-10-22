"""Builds the dashboard from small made-up segment files and checks the workbook and CSVs."""

import sys
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))
import dashboard  # noqa: E402
import segment  # noqa: E402

# Two runs. read_fields is the longest task step on average (10 s), so it is the bottleneck.
RUNS = {
    "run_01": [("open_invoice", 0, 4), ("read_fields", 4, 12), ("type_vendor", 12, 18), ("idle", 18, 40)],
    "run_02": [("open_invoice", 0, 2), ("read_fields", 2, 14), ("type_vendor", 14, 18),
               ("fix_error", 18, 22), ("submit", 22, 24)],
}


def make_outputs(tmp_path, with_evaluation=True):
    model_dir = tmp_path / "test-model"
    seg_rows, metric_rows = [], []
    for run, parts in RUNS.items():
        segs = [{"step": s, "start_s": a, "end_s": b} for s, a, b in parts]
        seg_rows += [(run, i, s["step"], s["start_s"], s["end_s"], s["end_s"] - s["start_s"])
                     for i, s in enumerate(segs, start=1)]
        metric_rows.append({"run": run, **segment.run_metrics(segs)})
    segment.write_csv(model_dir / "segments.csv", ["run", "order", "step", "start_s", "end_s", "duration_s"], seg_rows)
    segment.write_csv(model_dir / "run_metrics.csv", list(metric_rows[0].keys()), [list(m.values()) for m in metric_rows])
    if with_evaluation:
        pd.DataFrame([
            {"model": "test-model", "run": "run_01", "seconds_compared": 40, "frame_accuracy_raw": 0.8,
             "frame_accuracy": 0.85, "true_cycle_s": 40, "pred_cycle_s": 38, "cycle_abs_error_s": 2,
             "cycle_pct_error": 5, "step_duration_mae_s": 1.5, "order_distance": 0.1, "cost_usd": 0.1},
            {"model": "test-model", "run": "run_02", "seconds_compared": 24, "frame_accuracy_raw": 0.7,
             "frame_accuracy": 0.75, "true_cycle_s": 24, "pred_cycle_s": 24, "cycle_abs_error_s": 0,
             "cycle_pct_error": 0, "step_duration_mae_s": 2.5, "order_distance": 0.0, "cost_usd": ""},
        ]).to_csv(tmp_path / "metrics.csv", index=False)


def test_dashboard_workbook_and_powerbi_tables(tmp_path):
    make_outputs(tmp_path)
    workbook_path, powerbi_dir, steps, runs = dashboard.build_dashboard("test-model", tmp_path)

    wb = load_workbook(workbook_path)
    assert wb.sheetnames == ["Runs", "Steps", "Variation", "Evaluation", "About"]
    assert len(wb["Runs"]._charts) == 1
    assert len(wb["Steps"]._charts) == 1
    assert len(wb["Variation"]._charts) == 1

    bottleneck = steps[steps["is_bottleneck"]]
    assert list(bottleneck["step"]) == ["read_fields"]
    assert bottleneck.iloc[0]["mean_s"] == 10
    # idle has the largest mean (11 s) but is never the bottleneck
    assert steps.set_index("step").loc["idle", "mean_s"] == 11

    step_times = pd.read_csv(powerbi_dir / "step_times.csv")
    assert len(step_times) == 2 * len(dashboard.STEPS)            # every run x every step
    assert step_times.query("run == 'run_01' and step == 'fix_error'")["seconds"].item() == 0

    runs_csv = pd.read_csv(powerbi_dir / "runs.csv")
    assert list(runs_csv["cycle_s"]) == [40, 24]
    assert (powerbi_dir / "steps.csv").exists() and (powerbi_dir / "evaluation.csv").exists()

    evaluation = wb["Evaluation"]
    assert evaluation["A2"].value == "test-model" and evaluation["B2"].value == 2   # n runs
    assert "Bottleneck (highlighted): read_fields" in "".join(str(c.value) for c in wb["Steps"]["A"])


def test_dashboard_without_evaluation(tmp_path):
    make_outputs(tmp_path, with_evaluation=False)
    workbook_path, powerbi_dir, _, _ = dashboard.build_dashboard("test-model", tmp_path)
    assert "No evaluation yet" in load_workbook(workbook_path)["Evaluation"]["A1"].value
    assert not (powerbi_dir / "evaluation.csv").exists()
